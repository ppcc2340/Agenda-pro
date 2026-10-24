"""Exporta la base a un Excel con tres hojas: Respondieron, Enviados, Todos."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import store
from .config import RAIZ

COLUMNAS = [
    ("Nombre", "nombre", 32), ("Rubro", "rubro", 18), ("Ciudad", "ciudad", 20),
    ("Correo", "email", 30), ("Teléfono", "telefono", 16), ("Web", "web", 32),
    ("Instagram", "instagram", 30), ("Dirección", "direccion", 36), ("Fuente", "fuente", 10),
    ("Estado", "estado", 11), ("Plantilla", "plantilla", 20), ("Enviado", "enviado_el", 16), ("Respondió", "respondio_el", 16),
    ("Respuesta", "respuesta", 60),
]
COLOR_ESTADO = {"respondio": "C6EFCE", "enviado": "FFF2CC", "baja": "F4CCCC",
                "rebote": "E0E0E0", "nuevo": "DDEBF7", "sin_correo": "FFFFFF"}


def _hoja(wb, titulo, filas):
    ws = wb.create_sheet(titulo)
    ws.append([c[0] for c in COLUMNAS])
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F4E78")
        celda.alignment = Alignment(vertical="center")
    for f in filas:
        ws.append([f[c[1]] for c in COLUMNAS])
        color = COLOR_ESTADO.get(f["estado"], "FFFFFF")
        ws.cell(row=ws.max_row, column=10).fill = PatternFill("solid", fgColor=color)
    for i, (_, _, ancho) in enumerate(COLUMNAS, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def exportar(con, nombre_archivo: str) -> str:
    todos = store.listar(con)
    wb = Workbook()
    wb.remove(wb.active)
    _hoja(wb, "Respondieron", [f for f in todos if f["estado"] == "respondio"])
    _hoja(wb, "Enviados", [f for f in todos if f["estado"] in ("enviado", "respondio", "baja", "rebote")])
    _hoja(wb, "Todos los leads", todos)
    ruta = RAIZ / nombre_archivo
    wb.save(ruta)
    return str(ruta)
